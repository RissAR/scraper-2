import io
from typing import Any, Sequence

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def export_listings_to_excel(listings: Sequence[dict[str, Any]]) -> bytes:
    """
    Генерує розширений Excel (.xlsx) файл із переліком оголошень від власників
    з бази даних та повертає його байти.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Власники Дніпро"

    # Заголовки стовпців
    headers = [
        "ID",
        "Джерело",
        "Заголовок",
        "Ціна",
        "Валюта",
        "Телефон",
        "Власник (Ім'я)",
        "К-ть ог. в акаунта",
        "Район",
        "Поверх",
        "Площа",
        "Посилання (URL)",
        "Дата створення",
        "Реклама (ТОП)",
        "Додано в базу",
    ]

    ws.append(headers)

    # Стилізація заголовка
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Заповнення рядків
    for row_idx, item in enumerate(listings, start=2):
        created_val = str(item.get("created_time") or "")
        seen_val = str(item.get("seen_at") or "")
        is_promo = "Так" if item.get("is_promoted") else "Ні"
        floor_str = f"{item.get('floor', '')}/{item.get('total_floors', '')}" if item.get('floor') else ""

        row_data = [
            item.get("id", ""),
            item.get("source", "OLX.ua"),
            item.get("title", ""),
            item.get("price", 0.0),
            item.get("currency", "грн"),
            item.get("phone") or "Не вказано",
            item.get("author_name") or "Власник",
            item.get("author_ads_count", 1),
            item.get("district") or "Не вказано",
            floor_str,
            item.get("total_area") or "",
            item.get("url", ""),
            created_val[:19].replace("T", " ") if created_val else "",
            is_promo,
            seen_val[:19].replace("T", " ") if seen_val else "",
        ]
        ws.append(row_data)

        # Додаємо межі та вирівнювання
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if col_idx in (1, 2, 5, 8, 10, 14):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"

    # Автопідбір ширини колонок
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    ws.row_dimensions[1].height = 25

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
